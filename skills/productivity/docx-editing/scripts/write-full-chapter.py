#!/usr/bin/env python3
"""
Build a full academic Chapter 1 with citations and references from a Google Doc outline.

Usage:
  1. Export the Google Doc as docx (or download via curl)
  2. Edit the SRC path and content dicts below
  3. Run: python write-full-chapter.py

Pattern: for each section, define heading text + content paragraphs.
The script renumbers headings, inserts content, and appends a numbered
Danh Mục Tài Liệu Tham Khảo section.
"""
import docx
from copy import deepcopy

SRC = r"C:\Users\datel\Desktop\Chuong1_GoogleDoc.docx"
DST = r"C:\Users\datel\Desktop\Chương 1 - Hoàn chỉnh.docx"

doc = docx.Document(SRC)

def set_text(para, text):
    for run in para.runs:
        run.text = ""
    if para.runs:
        para.runs[0].text = text
    else:
        para.add_run(text)

def insert_multiple(ref_para, texts):
    current = ref_para._element
    body = doc.element.body
    for t in texts:
        new_p = deepcopy(current)
        for r in new_p.iterfind(
            ".//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}t"
        ):
            r.text = ""
        first_t = new_p.find(
            ".//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}t"
        )
        if first_t is not None:
            first_t.text = t
        idx = list(body).index(current)
        body.insert(idx + 1, new_p)
        current = new_p
    return current

# ── Step 1: Rename bare headings into numbered format ──
heading_map = {
    2: "1.1.1. Thông tin du lịch trên mạng xã hội",
    3: "1.1.2. Social Media Skepticism",
    4: "1.1.3. Hành vi kiểm chứng thông tin",
    5: "1.1.4. Hành vi sử dụng thông tin",
    7: "1.2.1. Information Credibility Theory",
    8: "1.2.2. Persuasion Knowledge Model",
    9: "1.2.3. Elaboration Likelihood Model (ELM)",
    10: "1.2.4. Social Cognitive Theory",
    12: "1.3.1. Quốc tế",
    13: "1.3.2. Việt Nam",
}
for idx, text in heading_map.items():
    set_text(doc.paragraphs[idx], text)

# ── Step 2: Insert content per section (edit these dicts) ──
sections = {}

# Example: insert after para 2 (1.1.1)
sections[2] = [
    "Content paragraph 1 here with citation [1].",
    "Content paragraph 2 with citation [2].",
]

for ref_idx, content_list in sections.items():
    insert_multiple(doc.paragraphs[ref_idx], content_list)

# ── Step 3: Add references section ──
# Find last real paragraph (skip sectPr)
body = doc.element.body
all_children = list(body)
last_elem = None
for elem in reversed(all_children):
    if not elem.tag.endswith("sectPr"):
        last_elem = elem
        break
last_para = None
for p in doc.paragraphs:
    if p._element is last_elem:
        last_para = p
        break

if last_para:
    refs = [
        "DANH MỤC TÀI LIỆU THAM KHẢO",
        "[1] Author, A. (Year). Title. Journal, Volume(Issue), Pages.",
    ]
    insert_multiple(last_para, refs)

doc.save(DST)
print(f"Done. Saved to {DST}")
